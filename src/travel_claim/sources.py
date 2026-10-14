"""Fresh native-source retrieval and byte-bound source manifests.

This module deliberately has no cache/fallback path. Every capture invokes the
configured native URL or its documented export route again.
"""

from __future__ import annotations

import hashlib
import io
from http.client import HTTPException, IncompleteRead
import json
import re
import urllib.error
import urllib.request
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Mapping


@dataclass(frozen=True)
class NativeSource:
    source_id: str
    title: str
    url: str
    locator: str
    export_url: str | None = None
    kind: str = "binary"


SOURCES: tuple[NativeSource, ...] = (
    NativeSource(
        "policy-pol-2026-2",
        "Travel reimbursement policy POL-2026.2",
        "https://private-pecorino-70e.notion.site/Project-D-Alderbridge-travel-reimbursement-policy-POL-2026-1-3d80b700541e81ed93ebd386abc1dcb0",
        "Notion policy page body",
        kind="notion",
    ),
    NativeSource(
        "initial-claim-binder",
        "Initial claim binder",
        "https://drive.google.com/file/d/13sRj6s_kMeP1hsYGstWWmfkb6SCN1YOS/view",
        "Complete PDF document",
        "https://drive.google.com/uc?export=download&id=13sRj6s_kMeP1hsYGstWWmfkb6SCN1YOS",
        "pdf",
    ),
    NativeSource(
        "claim-updates-binder",
        "Claim updates binder",
        "https://drive.google.com/file/d/16829RwrGrMCK9oq5TnWuytFewGk_8xQP/view",
        "Complete PDF document",
        "https://drive.google.com/uc?export=download&id=16829RwrGrMCK9oq5TnWuytFewGk_8xQP",
        "pdf",
    ),
    NativeSource(
        "receipt-binder",
        "Receipt binder",
        "https://drive.google.com/file/d/1Ynn1VLH-faRIy91xy95rOWwCBhy5mjjM/view",
        "Complete PDF document",
        "https://drive.google.com/uc?export=download&id=1Ynn1VLH-faRIy91xy95rOWwCBhy5mjjM",
        "pdf",
    ),
    NativeSource(
        "travel-registers-workbook",
        "Travel registers workbook",
        "https://docs.google.com/spreadsheets/d/1338nvKD6rgAdLgGRhYCMMlQyDNLK0sE8pvFWrZAaaxg/edit",
        "All workbook tabs via native XLSX export",
        "https://docs.google.com/spreadsheets/d/1338nvKD6rgAdLgGRhYCMMlQyDNLK0sE8pvFWrZAaaxg/export?format=xlsx",
        "xlsx",
    ),
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _policy_page_is_relevant_html(content: bytes) -> bool:
    text = content.decode("utf-8", "replace")
    # A successful response can still be only a generic Notion app/login shell.
    has_policy_id = bool(re.search(r"POL[- ]2026(?:\.2|[- ]2)", text, re.I))
    has_policy_body = bool(re.search(r"lodging|meal cap|permit approval|reimbursement", text, re.I))
    return has_policy_id and has_policy_body


def _validate_payload(source: NativeSource, content: bytes) -> tuple[bool, str | None, str | None]:
    """Return (usable, filename suffix, failure reason), without parsing case answers."""
    if source.kind == "notion":
        if not content or not _policy_page_is_relevant_html(content):
            return False, None, "HTTP response contained a generic Notion shell; policy page content was not accessible"
        return True, "html", None
    if source.kind == "pdf":
        if not content.startswith(b"%PDF-") or b"%%EOF" not in content[-2048:]:
            return False, None, "Export did not contain a complete PDF signature and trailer"
        return True, "pdf", None
    if source.kind == "xlsx":
        try:
            import io

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if archive.testzip() is not None or "xl/workbook.xml" not in archive.namelist():
                    return False, None, "XLSX archive is corrupt or has no workbook index"
        except (OSError, zipfile.BadZipFile):
            return False, None, "Export was not a readable XLSX workbook"
        return True, "xlsx", None
    return (bool(content), "bin" if content else None, None if content else "Empty response body")


def _request(url: str, timeout: float) -> tuple[bytes, Mapping[str, str], str]:
    req = urllib.request.Request(url, headers={"User-Agent": "travel-claim-source-capture/2.0", "Cache-Control": "no-cache", "Pragma": "no-cache"})
    with urllib.request.urlopen(req, timeout=timeout) as response:
        headers = dict(response.headers.items())
        headers["X-Travel-Claim-HTTP-Status"] = str(response.status)
        try:
            body = response.read()
        except IncompleteRead as exc:
            body = exc.partial
            headers["X-Travel-Claim-Read-Error"] = str(exc)
        return body, headers, response.geturl()


WORKBOOK_TABS = {"People", "Trips", "Merchant payments", "FX", "Caps", "Budget", "Review ledger",
                 "Review packets", "Finance activity", "Travel cancellations"}


def assess_payload(source: NativeSource, body: bytes) -> tuple[str, str | None, str | None, list[str]]:
    """Check accessible scope; a Notion HTML fragment never proves a full policy read."""
    usable, suffix, reason = _validate_payload(source, body)
    if not usable:
        return "unavailable", None, reason, [source.title]
    if source.kind == "notion":
        return "partial", suffix, "Policy-related HTML accessed, but complete native Notion page content is not established", ["Complete authoritative policy page body and all relevant blocks"]
    try:
        if source.kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(body), strict=True)
            if reader.is_encrypted or not reader.pages:
                raise ValueError("PDF has no accessible pages or is encrypted")
            # Traverse every page, not just the signature/trailer. Raw bytes are
            # retained even where text extraction cannot read scanned content.
            texts = [(page.extract_text(extraction_mode="layout") or "") if "/Contents" in page else "" for page in reader.pages]
            if any(not text.strip() for text in texts):
                return "partial", suffix, "One or more PDF pages have no extractable text; visual/OCR verification required", ["Readable text or verified visual transcription of every PDF page"]
        elif source.kind == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(body), data_only=False, read_only=False)
            missing = sorted(WORKBOOK_TABS - set(workbook.sheetnames))
            workbook.close()
            if missing:
                return "partial", suffix, "Workbook export lacks required retained-source tabs", missing
    except Exception as exc:
        return "unavailable", None, f"Native document validation failed: {type(exc).__name__}: {exc}", [source.title]
    return "read", suffix, None, []


def capture_sources(
    run_dir: Path, *, run_id: str, case_clock: str | None,
    source_specs: tuple[NativeSource, ...] = SOURCES,
    fetch: Callable[[str, float], tuple[bytes, Mapping[str, str], str]] = _request,
    observed_at: Callable[[], str] = _utc_now, timeout: float = 30.0,
) -> dict:
    """Attempt fresh requests; retain every returned body, including error evidence.

    A new, empty run directory is required. There is no retained-input option.
    Observation timestamps are retrieval facts, never the scenario case clock.
    """
    if not run_id or Path(run_id).name != run_id or run_id in {".", ".."}:
        raise ValueError("run_id must be a single safe path component")
    if run_dir.name != run_id:
        raise ValueError("run_dir basename must equal run_id")
    if timeout <= 0:
        raise ValueError("timeout must be positive")
    if case_clock is not None:
        try:
            clock = datetime.fromisoformat(case_clock.replace("Z", "+00:00"))
            if clock.utcoffset() is None:
                raise ValueError("timezone required")
        except (ValueError, AttributeError, TypeError) as exc:
            raise ValueError("case_clock must be an explicit date/time with timezone, never an observation default") from exc
    if set(source_specs) != set(SOURCES) or len(source_specs) != 5:
        raise ValueError("Every live capture must attempt exactly all five original source identities and URLs")
    if run_dir.exists() and any(run_dir.iterdir()):
        raise FileExistsError(f"Refusing to reuse nonempty run directory: {run_dir}")
    run_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = run_dir / "sources.json"
    (run_dir / "sources").mkdir()
    (run_dir / "responses").mkdir()
    entries = []
    for source in source_specs:
        started = observed_at()
        attempted = source.export_url or source.url
        body, headers, final_url, status_code, failure = None, {}, None, None, None
        try:
            body, headers, final_url = fetch(attempted, timeout)
            status_code = int(next((v for k, v in headers.items() if k.lower() == "x-travel-claim-http-status"), "200"))
        except urllib.error.HTTPError as exc:
            status_code, final_url = exc.code, exc.geturl()
            headers = dict(exc.headers.items()) if exc.headers else {}
            try:
                body = exc.read()
            except OSError:
                pass
            failure = f"HTTPError: {exc.code} {exc.reason}"
        except (urllib.error.URLError, TimeoutError, OSError, ValueError, HTTPException) as exc:
            failure = f"{type(exc).__name__}: {exc}"
        completed = observed_at()
        normalized = {k.lower(): v for k, v in headers.items()}
        selected_headers = {k: normalized[k] for k in ("etag", "last-modified", "content-type", "content-length", "content-range", "x-travel-claim-read-error") if k in normalized}
        response_path = None
        if body is not None:
            response_path = f"responses/{source.source_id}.body"
            with (run_dir / response_path).open("xb") as stream:
                stream.write(body)
        if failure:
            state, suffix, error, missing = "unavailable", None, failure, [source.title]
        else:
            state, suffix, error, missing = assess_payload(source, body or b"")
            if status_code != 200 or "content-range" in normalized:
                state, error, missing = "partial" if suffix else "unavailable", f"Incomplete HTTP response: status {status_code}, content range {normalized.get('content-range')}", [source.title]
            if "x-travel-claim-read-error" in normalized:
                state, error, missing = "partial" if suffix else "unavailable", normalized["x-travel-claim-read-error"], [source.title]
            if "content-length" in normalized:
                try:
                    if int(normalized["content-length"]) != len(body or b""):
                        state, error, missing = "partial" if suffix else "unavailable", "Response byte length differs from Content-Length", [source.title]
                except ValueError:
                    state, error, missing = "unavailable", "Invalid Content-Length", [source.title]
        rel_path = None
        if suffix and body is not None and state != "unavailable":
            rel_path = f"sources/{source.source_id}.{suffix}"
            with (run_dir / rel_path).open("xb") as stream:
                stream.write(body)
        entries.append({
            "source_id": source.source_id, "path": rel_path,
            "sha256": sha256_bytes(body) if rel_path else None,
            "status": state, "error": error, "url": source.url, "locator": source.locator,
            "observed_at": completed, "version": (normalized.get("etag") or normalized.get("last-modified")) if rel_path else None,
            "scope": source.title, "missing_scope": missing,
            "observation": {"method": "fresh-http-export" if source.export_url else "fresh-http-native",
                "requested_url": attempted, "final_url": final_url, "http_status": status_code,
                "started_at": started, "completed_at": completed, "response_headers": selected_headers,
                "response_path": response_path, "response_sha256": sha256_bytes(body) if body is not None else None,
                "response_bytes": len(body) if body is not None else None,
                "content_version": sha256_bytes(body) if body is not None else None,
                "cache_fallback_used": False},
        })
    manifest = {"run_id": run_id, "source_version": "native-source-manifest/2", "case_clock": case_clock, "files": entries}
    with manifest_path.open("x", encoding="utf-8", newline="\n") as stream:
        stream.write(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")
    return manifest


def verify_source_manifest(run_dir: Path) -> list[str]:
    """Check source manifest shape, five source attempts, and retained-byte hashes."""
    errors: list[str] = []
    path = run_dir / "sources.json"
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Cannot read sources.json: {exc}"]
    if not isinstance(manifest, dict):
        return ["sources.json must contain an object"]
    if set(manifest) != {"run_id", "source_version", "case_clock", "files"}:
        errors.append("sources.json must contain exactly run_id, source_version, case_clock, and files")
    if manifest.get("run_id") != run_dir.name:
        errors.append("sources.json run_id must match its run directory")
    if not isinstance(manifest.get("files"), list):
        return errors + ["sources.json files must be an array"]
    if manifest.get("source_version") not in {"native-source-manifest/1", "native-source-manifest/2"}:
        errors.append("Unsupported source manifest version")
    specs = {source.source_id: source for source in SOURCES}
    seen: set[str] = set()
    for record in manifest["files"]:
        if not isinstance(record, dict):
            errors.append("every files[] entry must be an object")
            continue
        sid = record.get("source_id")
        if not isinstance(sid, str) or not sid or sid in seen:
            errors.append(f"source_id is empty or duplicated: {sid!r}")
        else:
            seen.add(sid)
        required = {
            "source_id", "path", "sha256", "status", "error", "url", "locator",
            "observed_at", "version", "scope", "missing_scope",
        }
        if not required.issubset(record):
            errors.append(f"{sid}: missing required metadata fields {sorted(required - set(record))}")
            continue
        if not isinstance(record.get("url"), str) or not record["url"] or not isinstance(record.get("locator"), str) or not record["locator"]:
            errors.append(f"{sid}: url and locator must be nonempty strings")
        if not isinstance(record.get("scope"), str) or not record["scope"]:
            errors.append(f"{sid}: scope must identify the observed source portion")
        if not isinstance(record.get("missing_scope"), list) or not all(isinstance(x, str) and x for x in record["missing_scope"]):
            errors.append(f"{sid}: missing_scope must be an array of nonempty strings")
        try:
            observed = datetime.fromisoformat(str(record["observed_at"]).replace("Z", "+00:00"))
            if observed.utcoffset() is None or observed.utcoffset().total_seconds() != 0:
                errors.append(f"{sid}: observed_at must be UTC")
        except (TypeError, ValueError):
            errors.append(f"{sid}: observed_at must be an ISO UTC timestamp")
        if record["error"] is not None and not isinstance(record["error"], str):
            errors.append(f"{sid}: error must be a string or null")
        if sid in specs and record["url"] != specs[sid].url:
            errors.append(f"{sid}: source identity does not match its original native URL")
        if manifest.get("source_version") == "native-source-manifest/2":
            errors.extend(_verify_observation(run_dir, record, specs.get(sid)))
        if record["status"] in {"read", "partial"}:
            if record["status"] == "read" and (record["missing_scope"] or record["error"]):
                errors.append(f"{sid}: read scope cannot contain missing scope or errors")
            if record["status"] == "partial" and (not record["missing_scope"] or not record["error"]):
                errors.append(f"{sid}: partial scope must identify missing content and an error")
            if not isinstance(record["path"], str) or not record["path"] or not isinstance(record["sha256"], str) or not record["sha256"]:
                errors.append(f"{sid}: read record has no retained path or hash")
                continue
            content_path = run_dir / record["path"]
            if not content_path.resolve().is_relative_to(run_dir.resolve()):
                errors.append(f"{sid}: retained path escapes the run directory")
                continue
            try:
                actual = sha256_bytes(content_path.read_bytes())
            except OSError as exc:
                errors.append(f"{sid}: cannot read retained content: {exc}")
            else:
                if actual != record["sha256"]:
                    errors.append(f"{sid}: retained content hash mismatch")
        elif record["status"] == "unavailable":
            if any(record[k] is not None for k in ("path", "sha256", "version")):
                errors.append(f"{sid}: unavailable scope must have null path, sha256, and version")
            if not record["error"] or not record["missing_scope"]:
                errors.append(f"{sid}: unavailable scope has no error or missing scope")
        else:
            errors.append(f"{sid}: invalid status {record['status']!r}")
    if seen != set(specs) or len(manifest["files"]) != 5:
        errors.append("manifest must cover exactly all five configured native source identities")
    expected_urls = {source.url for source in SOURCES}
    observed_urls = {
        record.get("url")
        for record in manifest["files"]
        if isinstance(record, dict) and isinstance(record.get("url"), str)
    }
    if observed_urls != expected_urls:
        errors.append("manifest must cover exactly all five configured native source URLs")
    return errors


def _verify_observation(run_dir: Path, record: dict, source: NativeSource | None) -> list[str]:
    errors = []
    sid = record.get("source_id")
    observation = record.get("observation")
    if not source or record.get("url") != source.url or not isinstance(observation, dict):
        return [f"{sid}: native identity, URL or fresh retrieval observation missing"]
    if (observation.get("requested_url") != (source.export_url or source.url)
            or observation.get("method") != ("fresh-http-export" if source.export_url else "fresh-http-native")
            or observation.get("cache_fallback_used") is not False):
        errors.append(f"{sid}: fresh native request binding or no-fallback evidence invalid")
    try:
        start, end = [datetime.fromisoformat(str(observation[k]).replace("Z", "+00:00")) for k in ("started_at", "completed_at")]
        if any(t.utcoffset() is None or t.utcoffset().total_seconds() != 0 for t in (start, end)) or end < start or observation['completed_at'] != record['observed_at']:
            raise ValueError("retrieval timestamps must be ordered UTC observations")
    except (ValueError, TypeError, KeyError):
        errors.append(f"{sid}: invalid retrieval observation times")
    path = observation.get("response_path")
    if path is None:
        if record.get("status") != "unavailable" or any(observation.get(k) is not None for k in ("response_sha256", "response_bytes", "content_version")):
            errors.append(f"{sid}: missing response cannot establish accessible content")
        return errors
    try:
        response_path = run_dir / path
        if not response_path.resolve().is_relative_to(run_dir.resolve()):
            raise ValueError("response path escapes run directory")
        body = response_path.read_bytes()
        digest = sha256_bytes(body)
        if digest != observation.get("response_sha256") or digest != observation.get("content_version") or len(body) != observation.get("response_bytes"):
            errors.append(f"{sid}: response evidence hash or size mismatch")
        if record.get("status") in {"read", "partial"} and record.get("sha256") != digest:
            errors.append(f"{sid}: retained source differs from fresh response evidence")
        if record.get("status") == "read":
            state, _, reason, _ = assess_payload(source, body)
            headers = observation.get("response_headers", {})
            if state != "read" or observation.get("http_status") != 200 or "content-range" in headers or "x-travel-claim-read-error" in headers:
                errors.append(f"{sid}: complete read not supported by native response: {reason}")
            if "content-length" in headers and int(headers['content-length']) != len(body):
                errors.append(f"{sid}: complete read has inconsistent Content-Length")
    except (OSError, ValueError, TypeError) as exc:
        errors.append(f"{sid}: cannot verify response evidence: {exc}")
    return errors
