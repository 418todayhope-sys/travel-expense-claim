"""Fresh source requests never borrow retained inputs or assert missing content."""
import hashlib
import io
import json
from pathlib import Path
from urllib.error import HTTPError, URLError

import pytest

from source_fixtures import pdf_bytes, xlsx_bytes
from travel_claim.cli import main
from travel_claim.live_run import run_live_sources
from travel_claim.parsing import parse_retained_run
from travel_claim.sources import SOURCES, capture_sources, verify_source_manifest


def fetcher(*, fault=None, calls=None):
    def fetch(url, timeout):
        if calls is not None: calls.append(url)
        source = next(s for s in SOURCES if (s.export_url or s.url) == url)
        if fault == 'all-fail': raise URLError('SYNTHETIC network unavailable')
        if source.kind == 'notion':
            if fault == 'http-error':
                raise HTTPError(url, 403, 'SYNTHETIC denied', {'ETag': 'denied-v1'}, io.BytesIO(b'SYNTHETIC access denied'))
            return b'<article>POL-2026.2 reimbursement SYNTHETIC fragment</article>', {}, url
        body = xlsx_bytes() if source.kind == 'xlsx' else pdf_bytes()
        headers = {'etag': 'synthetic-version', 'content-length': str(len(body))}
        if source.source_id == 'initial-claim-binder':
            if fault == 'range': headers.update({'X-Travel-Claim-HTTP-Status': '206', 'Content-Range': 'bytes 0-9/100'})
            elif fault == 'truncated': headers['content-length'] = str(len(body) + 100)
            elif fault == 'corrupt': body = b'%PDF-1.4 broken %%EOF'
        return body, headers, url + '#SYNTHETIC-final-url'
    return fetch


def test_successful_native_document_retrieval_preserves_bytes_versions_and_all_five_requests(tmp_path):
    calls = []
    run = tmp_path / 'fresh-a'
    m = capture_sources(run, run_id=run.name, case_clock=None, fetch=fetcher(calls=calls))
    assert calls == [s.export_url or s.url for s in SOURCES]
    assert [r['status'] for r in m['files']] == ['partial', 'read', 'read', 'read', 'read']
    for row in m['files']:
        observation = row['observation']
        assert observation['cache_fallback_used'] is False
        body = (run / observation['response_path']).read_bytes()
        assert hashlib.sha256(body).hexdigest() == observation['response_sha256'] == row['sha256']
        assert body == (run / row['path']).read_bytes()
    assert m['files'][1]['version'] == 'synthetic-version'
    assert m['case_clock'] is None and verify_source_manifest(run) == []


@pytest.mark.parametrize('fault', ['range', 'truncated', 'corrupt', 'all-fail', 'http-error'])
def test_partial_or_failed_live_run_never_reports_complete_success(tmp_path, fault):
    run = tmp_path / fault
    report = run_live_sources(run, run_id=run.name, case_clock=None, fetch=fetcher(fault=fault))
    assert report['status'] == report['acquisition_status'] == 'BLOCKED'
    assert report['final_live_source_compliance'] is False and report['batches_consumed'] == []
    assert report['fresh_reads_attempted'] == 5
    assert not (run / 'batches').exists()
    assert verify_source_manifest(run) == []
    assert 'BLOCKED' in (run / 'live-run.md').read_text()
    if fault == 'http-error':
        row = json.loads((run / 'sources.json').read_text())['files'][0]
        assert row['observation']['http_status'] == 403
        assert (run / row['observation']['response_path']).read_bytes() == b'SYNTHETIC access denied'
    if fault in {'range', 'truncated'}:
        m = json.loads((run / 'sources.json').read_text())
        assert m['files'][1]['status'] == 'partial'
        assert 'initial-claim-binder' in {s['source_id'] for s in json.loads((run / 'parsed-records.json').read_text())['source_inventory']}


def test_failed_new_run_does_not_fall_back_to_successful_prior_inputs(tmp_path):
    first, second = tmp_path / 'prior', tmp_path / 'new'
    capture_sources(first, run_id=first.name, case_clock=None, fetch=fetcher())
    before = {str(p): p.read_bytes() for p in first.rglob('*') if p.is_file()}
    calls = []
    run_live_sources(second, run_id=second.name, case_clock=None, fetch=fetcher(fault='all-fail', calls=calls))
    assert len(calls) == 5
    m = json.loads((second / 'sources.json').read_text())
    assert all(r['status'] == 'unavailable' and r['path'] is None for r in m['files'])
    assert all(Path(p).read_bytes() == data for p, data in before.items())
    with pytest.raises(FileExistsError):
        capture_sources(first, run_id=first.name, case_clock=None, fetch=lambda *_: pytest.fail('must not request'))


def test_preexisting_cached_inputs_and_incomplete_source_list_are_rejected(tmp_path):
    run = tmp_path / 'cached'; run.mkdir(); (run / 'parsed-records.json').write_text('SYNTHETIC stale')
    with pytest.raises(FileExistsError):
        run_live_sources(run, run_id=run.name, case_clock=None, fetch=lambda *_: pytest.fail('must reject first'))
    with pytest.raises(ValueError):
        capture_sources(tmp_path / 'subset', run_id='subset', case_clock=None, source_specs=SOURCES[1:])


def test_forged_complete_policy_and_response_tampering_fail_verification(tmp_path):
    run = tmp_path / 'tampered'
    m = capture_sources(run, run_id=run.name, case_clock=None, fetch=fetcher())
    m['files'][0].update(status='read', error=None, missing_scope=[])
    (run / 'sources.json').write_text(json.dumps(m))
    assert any('complete read not supported' in e for e in verify_source_manifest(run))
    (run / m['files'][1]['observation']['response_path']).write_bytes(b'changed')
    assert any('response evidence hash' in e for e in verify_source_manifest(run))


def test_cli_live_and_capture_partial_return_blocked_and_no_clock_default(tmp_path, monkeypatch, capsys):
    import travel_claim.cli as cli
    original = cli.run_live_sources
    monkeypatch.setattr(cli, 'run_live_sources', lambda *a, **kw: original(*a, **kw, fetch=fetcher()))
    assert main(['live-run', 'cli-live', str(tmp_path / 'cli-live')]) == 2
    assert 'BLOCKED' in capsys.readouterr().out
    monkeypatch.setattr(cli, 'capture_sources', lambda *a, **kw: capture_sources(*a, **kw, fetch=fetcher()))
    assert main(['capture-sources', 'cli-capture', str(tmp_path / 'cli-capture')]) == 2
    assert 'partial' in capsys.readouterr().out


def test_each_new_run_makes_five_new_calls_even_if_versions_and_bytes_match(tmp_path):
    calls=[]
    for name in ('fresh-one', 'fresh-two'):
        capture_sources(tmp_path/name, run_id=name, case_clock=None, fetch=fetcher(calls=calls))
    assert len(calls) == 10 and calls[:5] == calls[5:]


def test_missing_workbook_scope_and_scanned_pdf_are_partial(tmp_path):
    from openpyxl import Workbook
    from pypdf import PdfWriter
    workbook = Workbook(); workbook.active.title = 'Trips'
    xlsx = io.BytesIO(); workbook.save(xlsx); workbook.close()
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    pdf = io.BytesIO(); writer.write(pdf)
    ordinary = fetcher()
    def fetch(url, timeout):
        if url == SOURCES[-1].export_url: return xlsx.getvalue(), {}, url
        if url == SOURCES[1].export_url: return pdf.getvalue(), {}, url
        return ordinary(url, timeout)
    run = tmp_path/'partial-documents'
    manifest = capture_sources(run, run_id=run.name, case_clock=None, fetch=fetch)
    assert manifest['files'][1]['status'] == manifest['files'][-1]['status'] == 'partial'
    assert 'People' in manifest['files'][-1]['missing_scope']
    assert verify_source_manifest(run) == []
    parsed = parse_retained_run(run)
    assert len(parsed['pdf_pages']) == 3  # partial page preserved alongside readable documents
    assert parsed['pdf_pages'][0]['text'] == ''
    assert parsed['missing_sources']


@pytest.mark.parametrize('clock', ['2026-09-01', '2026-09-01T12:00:00', 'invented-default'])
def test_incomplete_clock_values_are_rejected_before_capture(tmp_path, clock):
    with pytest.raises(ValueError, match='timezone'):
        run_live_sources(tmp_path/'bad-clock', run_id='bad-clock', case_clock=clock,
                         fetch=lambda *_: pytest.fail('must reject before fetching'))


def test_parser_preserves_fractional_batch_as_an_issue_instead_of_rounding(tmp_path):
    run=tmp_path/'fractional-batch'
    manifest=capture_sources(run, run_id=run.name, case_clock=None, fetch=fetcher())
    from openpyxl import load_workbook
    source=manifest['files'][-1]
    path=run/source['path']
    workbook=load_workbook(path)
    sheet=workbook['Finance activity']
    sheet.append(['Activity reference', 'Arrival batch', 'Activity type', 'Amount (EUR)'])
    sheet.append(['SYN-EVENT', 1.5, 'settled', 12.34])
    workbook.save(path); workbook.close()
    # Deliberately reconstruct a legacy synthetic parser fixture, not a fresh observation.
    manifest['source_version']='native-source-manifest/1'
    source['sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
    (run/'sources.json').write_text(json.dumps(manifest))
    parsed=parse_retained_run(run)
    assert parsed['arrival_batches']==[]
    value=parsed['records']['finance_events'][0]['values']['Arrival batch']['value']
    assert value==1.5
    assert any('not an explicit positive integer' in issue['error'] for issue in parsed['transcription_issues'])


def test_transport_uses_unconditional_cache_revalidation_and_preserves_interrupted_body(monkeypatch):
    from http.client import IncompleteRead
    from travel_claim.sources import _request
    requests=[]
    class Response:
        status=200
        headers={'Content-Length':'100'}
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self): raise IncompleteRead(b'SYNTHETIC partial', 100)
        def geturl(self): return SOURCES[0].url
    def open_request(request, timeout):
        requests.append(request)
        return Response()
    monkeypatch.setattr('urllib.request.urlopen', open_request)
    body, headers, url=_request(SOURCES[0].url, 1)
    assert body==b'SYNTHETIC partial' and 'X-Travel-Claim-Read-Error' in headers
    assert requests[0].get_header('Cache-control')=='no-cache'
    assert requests[0].get_header('If-none-match') is None
