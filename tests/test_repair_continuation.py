"""End-to-end synthetic employee evidence; no real responses or authority."""
from copy import deepcopy
import csv
import io
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from test_receipt_continuation import base_run, pdf, receipt, sha, hashes, SCHEMA, POLICY
from travel_claim.cli import main
from travel_claim.ordinary import fields, ref
from travel_claim.parsing import parse_retained_run, _workbook_sheets, _pdf_pages, _receipt_records
from travel_claim.policy_evidence import create_policy_replay
from travel_claim.receipt_continuation import verify_continuation
from travel_claim.sources import SOURCES, capture_sources


def get(path):
    return json.loads(path.read_text())


def save(path, data):
    path.write_text(json.dumps(data, indent=2))


def parent(tmp_path, *, missing_merchants=('A',), returned=False, corrected=False):
    seed = tmp_path / 'seed'; seed.mkdir()
    first = base_run(seed)
    workbook = load_workbook(first / 'sources/travel-registers-workbook.xlsx')
    tab = workbook['Merchant payments']
    for row in reversed(range(3, tab.max_row + 1)):
        if tab.cell(row, 2).value in {'COST-' + c for c in missing_merchants}:
            tab.delete_rows(row)
    if returned:
        original = get(first / 'parsed-records.json')
        sample = get(Path(__file__).parent / 'fixtures/ordinary_claim_c01.json')
        packet = {k: v['value'] for k, v in next(s for s in sample['workbook_sheets'] if s['name']=='Review packets')['records'][0]['values'].items()}
        packet.update({'Decision reference':'SYN-RETURN-A', 'Trip reference':'T-A', 'Expense references':'COST-A'})
        decision = {k:v['value'] for k,v in sample['records']['decisions'][0]['values'].items()}
        decision.update({'Decision reference':'SYN-RETURN-A','Subject reference':'SYN-A', 'Outcome':'return',
                         'Repair requested':'Supply receipt and settled merchant proof for COST-A',
                         'Source revisions':'POL-2026.2\n'+sha(first/'sources/initial-claim-binder.pdf')})
        for name, row in [('Review packets',packet),('Review ledger',decision)]:
            sheet=workbook[name];sheet.delete_rows(1,sheet.max_row)
            sheet.append(list(row));sheet.append(list(row.values()))
    buffer=io.BytesIO();workbook.save(buffer)
    bodies={s.source_id: (first/next(r['path'] for r in get(first/'sources.json')['files']
                                   if r['source_id']==s.source_id)).read_bytes()
            for s in SOURCES if s.source_id!='policy-pol-2026-2'}
    bodies.update({'policy-pol-2026-2':b'SYNTHETIC generic Notion shell', 'travel-registers-workbook':buffer.getvalue()})
    if corrected:
        from travel_claim.parsing import _pdf_pages
        original_page=_pdf_pages(first/'sources/initial-claim-binder.pdf','synthetic','synthetic')[0]['text']
        bodies['claim-updates-binder']=pdf([original_page.replace('Revision: 1','Revision: 2').replace('Arrival batch: 1','Arrival batch: 3')])
    def fetch(url, timeout):
        s=next(s for s in SOURCES if (s.export_url or s.url)==url)
        return bodies[s.source_id],{},url
    source=tmp_path/'source'
    capture_sources(source,run_id='source',case_clock=None,fetch=fetch,observed_at=lambda:'2026-10-01T12:00:00Z')
    parse_retained_run(source)
    result=tmp_path/'parent'
    create_policy_replay(source,result,run_id='parent',browser_policy=POLICY,schema_path=SCHEMA)
    return result


def bundle(tmp_path, parent_path, *, name='one', specs=None, insufficient=()):
    """spec = (claim suffix, item, kind, superseded refs)."""
    incoming=tmp_path/('incoming-'+name);incoming.mkdir()
    rows=[];outcomes={};assessments={}
    specs=specs or [('A','ordinary:SYN-A:r1:supporting_evidence','receipt',[])]
    for index,(cid,item,kind,supersedes) in enumerate(specs):
        sid=f'SYN-{name}-{index}';path=incoming/(sid+('.pdf' if kind=='receipt' else '.xlsx'))
        if kind=='receipt':
            path.write_bytes(pdf([receipt(cid)]));record=_receipt_records(_pdf_pages(path,sid,sha(path)))[0]
        else:
            workbook=Workbook();sheet=workbook.active;sheet.title='Merchant payments'
            value={'Transaction reference':'SYN-M-'+cid,'Expense reference':'COST-'+cid,'Employee ID':'EMP-01',
                   'Trip reference':'T-'+cid,'Currency':'EUR','Gross amount':'100.00','Paid date':'2026-09-01',
                   'Payment status':'settled','Receipt reference':'R-COST-'+cid}
            sheet.append(list(value));sheet.append(list(value.values()));workbook.save(path)
            record=_workbook_sheets(path,sid,sha(path))[0]['records'][0]
        rows.append(dict(source_id=sid,path=path.name,sha256=sha(path),status='read',error=None,missing_scope=[],
                         version=sha(path),observed_at='2026-10-02T12:00:00Z'))
        outcome=outcomes.setdefault(item,dict(outcome='supported',note='SYNTHETIC response for '+item,source_refs=[],repairs=[]))
        outcome['source_refs'].append(ref(record));outcome['repairs'].append(dict(kind=kind,record_ref=ref(record),supersedes_refs=supersedes))
    for index,(item,outcome) in enumerate(outcomes.items()):
        cid=item.split(':')[1]
        revision=int(item.split(':')[2][1:])
        assessments[str(index)]=dict(item_id=item,subject_type='claim',subject_id=cid,revision=revision,response_id='SYN-'+name,
            actor_id='EMP-01',sufficient=True,reason='SYNTHETIC supplied evidence for exact expense',
            provenance='SYNTHETIC employee response',evidence_refs=outcome['source_refs'],repairs=outcome['repairs'])
        if ':review_return:' in item:
            state=get(parent_path/('offline-replay-state.json' if (parent_path/'offline-replay-state.json').exists() else 'continuation-state.json'))
            from travel_claim.repair_continuation import return_context
            decision=return_context(state,item,cid,revision)
            assessments[str(index)].update(return_decision_ref=ref(decision['raw_record']),repair_requested=decision['repair_requested'])
    proof=incoming/'assessment.json';save(proof,assessments)
    proof_id='SYN-assess-'+name
    rows.append(dict(source_id=proof_id,path=proof.name,sha256=sha(proof),status='read',error=None,missing_scope=[],
                     version=sha(proof),observed_at='2026-10-02T12:00:00Z'))
    for index,outcome in enumerate(outcomes.values()):outcome['assessment_ref']=f'{proof_id}@{sha(proof)}#/{index}'
    for item in insufficient:outcomes[item]=dict(outcome='insufficient',note='SYNTHETIC still missing',source_refs=[])
    save(incoming/'sources.json',dict(files=rows))
    response=tmp_path/('response-'+name+'.json')
    save(response,dict(response_version='repair-response/1',response_id='SYN-'+name,actor_id='EMP-01',item_outcomes=outcomes))
    return response,incoming


def run(parent_path, response, evidence, output):
    return main(['continue-repairs',str(parent_path),str(response),str(output),'--evidence-dir',str(evidence),'--run-id',output.name])


def stage(state,cid,name):
    return next(s for s in state['ordinary_workflows'][cid]['stages'] if s['stage']==name)['status']


def test_multiple_evidence_types_partial_response_and_repeated_resumption(tmp_path):
    p=parent(tmp_path);before=hashes(p);old=get(p/'offline-replay-state.json')
    a='ordinary:SYN-A:r1:supporting_evidence';b='ordinary:SYN-B:r1:supporting_evidence'
    response,evidence=bundle(tmp_path,p,insufficient=[b]);one=tmp_path/'one'
    assert run(p,response,evidence,one)==0
    first=get(one/'continuation-state.json')
    assert stage(first,'SYN-A','supporting_evidence')=='BLOCKED'  # Receipt alone cannot prove payment.
    assert first['repair_queue']['items'][a]['status']=='open'
    assert first['repair_queue']['items'][b]['status']=='open'
    second_response,second_evidence=bundle(tmp_path,one,name='two',specs=[('A',a,'merchant-payment',[])],insufficient=[b])
    two=tmp_path/'two';assert run(one,second_response,second_evidence,two)==0
    final=get(two/'continuation-state.json');verify_continuation(two)
    assert stage(final,'SYN-A','supporting_evidence')=='PASS'
    assert final['repair_queue']['items'][a]['status']=='resolved'
    assert final['repair_queue']['items'][b]['status']=='open'
    assert final['ordinary_workflows']['SYN-B']==old['ordinary_workflows']['SYN-B']
    assert final['ordinary_workflows']['SYN-C']==old['ordinary_workflows']['SYN-C']
    assert final['finance_reconciliations']==old['finance_reconciliations']
    assert final['decision_history']==old['decision_history']
    assert final['revision_history']==old['revision_history']
    assert final['cancellation_state']==old['cancellation_state']
    assert final['budget_reconciliation']['scopes']==old['budget_reconciliation']['scopes']
    assert final['budget_history']==old['budget_history']
    assert final['ordinary_workflows']['SYN-A']['status']=='pending'
    assert final['ordinary_workflows']['SYN-A']['request'] is None
    assert all(not w['status'].startswith('closed-') for w in final['ordinary_workflows'].values())
    drafts=(two/'follow-up-drafts.md').read_text();assert f'Item `{a}`' not in drafts and f'Item `{b}`' in drafts
    assert 'supporting_evidence=PASS' in (two/'report.md').read_text()
    assert (two/final['native_source_manifest_path']).resolve()==(p/'sources.json').resolve()
    row=next(r for r in csv.DictReader((two/'claims.csv').open()) if r['claim_id']=='SYN-A')
    assert row['status']=='pending' and row['paid_cents']=='10000' and row['balance_cents']==''
    assert row['revision']=='1' and row['trip_id']=='T-A'
    assert row['allowed_cents']==str(final['ordinary_workflows']['SYN-A']['allowed_cents'])
    unchanged=hashes(two);assert run(one,second_response,second_evidence,two)==0 and hashes(two)==unchanged
    # Same response redelivered against the advanced state creates no new effect/result.
    assert run(two,second_response,second_evidence,tmp_path/'retry')==0
    assert not (tmp_path/'retry').exists() and hashes(two)==unchanged
    another,another_evidence=bundle(tmp_path,two,name='three',specs=[('B',b,'receipt',[])])
    assert run(two,another,another_evidence,tmp_path/'reuse'/'one')==1
    assert not (tmp_path/'reuse').exists()
    assert hashes(p)==before and len(list((p/'batches').glob('*.json')))==4
    for key,item in old['repair_queue']['items'].items():
        if item['subject_id'] not in {'SYN-A','SYN-B'}:assert final['repair_queue']['items'][key]==item


def test_valid_item_progresses_while_other_invalid_evidence_is_held(tmp_path):
    p=parent(tmp_path,missing_merchants=())
    a='ordinary:SYN-A:r1:supporting_evidence';b='ordinary:SYN-B:r1:supporting_evidence'
    response,evidence=bundle(tmp_path,p,specs=[('A',a,'receipt',[]),('B',b,'receipt',[])])
    bad=evidence/'SYN-one-1.pdf';bad.write_bytes(bad.read_bytes()+b'tampered')
    out=tmp_path/'mixed';assert run(p,response,evidence,out)==0
    state=get(out/'continuation-state.json');q=state['repair_queue']
    assert q['items'][a]['status']=='resolved' and q['items'][b]['status']=='open'
    event=next(e for e in q['events'] if e.get('item_id')==b and e['event']=='response-recorded')
    assert event['admission_status']=='held' and 'hash' in event['admission_reason']
    assert (out/'response-evidence'/next(r['path'] for r in get(out/'response-evidence/sources.json')['files'] if r['source_id']=='SYN-one-1')).read_bytes()==bad.read_bytes()
    verify_continuation(out)


@pytest.mark.parametrize('fault',['stale','wrong-actor','wrong-hash','missing-file','wrong-revision-assessment','wrong-operation','unknown-ref','future','wrong-gross','invalid-number','unsupported'])
def test_negative_deliveries_never_resolve_or_authorize(tmp_path,fault):
    p=parent(tmp_path);before=hashes(p);response,evidence=bundle(tmp_path,p)
    obj=get(response);key=next(iter(obj['item_outcomes']))
    if fault=='stale':obj['item_outcomes'][key.replace(':r1:',':r2:')]=obj['item_outcomes'].pop(key)
    elif fault=='wrong-actor':obj['actor_id']='UNAUTHORIZED'
    elif fault=='wrong-hash':(evidence/'SYN-one-0.pdf').write_bytes(b'tampered')
    elif fault=='missing-file':(evidence/'SYN-one-0.pdf').unlink()
    elif fault=='unknown-ref':obj['item_outcomes'][key]['source_refs']=['missing@'+'0'*64+'#row 1']
    elif fault=='unsupported':obj['item_outcomes']['ordinary:SYN-A:r1:required_approvals']=obj['item_outcomes'].pop(key)
    elif fault=='future':
        manifest=get(evidence/'sources.json');manifest['files'][0]['observed_at']='2026-11-16T12:00:00Z';save(evidence/'sources.json',manifest)
    else:
        proof=get(evidence/'assessment.json')
        if fault=='wrong-revision-assessment':proof['0']['revision']=2
        elif fault=='wrong-operation':proof['0']['repairs'][0]['kind']='approval'
        else:
            # A real, hash-bound PDF with mismatched gross is not sufficient either.
            path=evidence/'SYN-one-0.pdf';path.write_bytes(pdf([receipt(amount='invalid' if fault=='invalid-number' else '99.00')]))
            rr=f'SYN-one-0@{sha(path)}#PDF page 1, receipt 1'
            obj['item_outcomes'][key]['source_refs']=[rr];obj['item_outcomes'][key]['repairs'][0]['record_ref']=rr
            proof['0']['evidence_refs']=[rr];proof['0']['repairs'][0]['record_ref']=rr
        save(evidence/'assessment.json',proof);manifest=get(evidence/'sources.json')
        for row in manifest['files']:row['sha256']=sha(evidence/row['path'])
        save(evidence/'sources.json',manifest);obj['item_outcomes'][key]['assessment_ref']=f'SYN-assess-one@{sha(evidence/"assessment.json")}#/0'
    save(response,obj);out=tmp_path/'negative'
    rc=run(p,response,evidence,out)
    if fault in {'stale','wrong-actor'}:assert rc==1 and not out.exists()
    else:
        assert rc==0
        state=get(out/'continuation-state.json');assert state['response_holds']
        assert all(i['status']=='open' for k,i in state['repair_queue']['items'].items() if k in obj['item_outcomes'])
        assert state['finance_reconciliations']==get(p/'offline-replay-state.json')['finance_reconciliations']
    assert hashes(p)==before


def test_returned_work_receives_evidence_but_requires_human_re_review(tmp_path):
    p=parent(tmp_path,returned=True)
    old=get(p/'offline-replay-state.json');wf=old['ordinary_workflows']['SYN-A']
    item='ordinary:SYN-A:r1:review_return:SYN-RETURN-A'
    assert item in old['repair_queue']['items'] and old['repair_queue']['items'][item]['status']=='open'
    response,evidence=bundle(tmp_path,p,specs=[('A',item,'receipt',[])])
    partial=tmp_path/'returned-partial';assert run(p,response,evidence,partial)==0
    assert get(partial/'continuation-state.json')['repair_queue']['items'][item]['status']=='open'
    response2,evidence2=bundle(tmp_path,partial,name='returned-two',specs=[('A',item,'merchant-payment',[])])
    out=tmp_path/'returned';assert run(partial,response2,evidence2,out)==0
    state=get(out/'continuation-state.json')
    assert stage(state,'SYN-A','supporting_evidence')=='PASS'
    assert stage(state,'SYN-A','required_approvals')=='BLOCKED'
    assert state['repair_queue']['items'][item]['status']=='open'
    assert state['repair_queue']['items'][item]['owner']=='ADMIN-01'
    assert state['ordinary_workflows']['SYN-A']['status']=='pending'
    assert state['decision_history']==old['decision_history']
    assert state['finance_reconciliations']==old['finance_reconciliations']
    assert 'human re-review' in (out/'report.md').read_text()


def test_original_evidence_and_derived_result_tampering_are_detected(tmp_path):
    p=parent(tmp_path);response,evidence=bundle(tmp_path,p);one=tmp_path/'one'
    assert run(p,response,evidence,one)==0
    (one/'claims.csv').write_text('tampered')
    assert run(p,response,evidence,one)==1


def test_explicit_replacement_preserves_old_rows_and_conflicting_replay_is_rejected(tmp_path):
    # Explicitly replace merchant proof, without changing any Finance event.
    p=parent(tmp_path,missing_merchants=())
    key='ordinary:SYN-A:r1:supporting_evidence'
    parsed=get(p/'parsed-records.json')
    old_payment=next(r for s in parsed['workbook_sheets'] if s['name']=='Merchant payments'
                     for r in s['records'] if fields(r).get('Expense reference')=='COST-A')
    response,evidence=bundle(tmp_path,p,specs=[('A',key,'receipt',[]),('A',key,'merchant-payment',[ref(old_payment)])])
    out=tmp_path/'replaced';assert run(p,response,evidence,out)==0
    state=get(out/'continuation-state.json')
    assert stage(state,'SYN-A','supporting_evidence')=='PASS'
    changed=state['repair_deliveries'][0]['admitted'][0]['changes']
    assert changed[1]['superseded']==[old_payment]
    assert ref(old_payment) in get(response)['item_outcomes'][key]['repairs'][1]['supersedes_refs']
    assert state['finance_reconciliations']==get(p/'offline-replay-state.json')['finance_reconciliations']
    before=hashes(out);obj=get(response);obj['item_outcomes'][key]['note']='conflicting replay';save(response,obj)
    assert run(out,response,evidence,tmp_path/'conflict')==1 and hashes(out)==before


def test_unknown_supersession_is_held_and_unmatched_work_is_untouched(tmp_path):
    p=parent(tmp_path)
    key='ordinary:SYN-A:r1:supporting_evidence'
    response,evidence=bundle(tmp_path,p,specs=[('A',key,'receipt',['invented-ref'])])
    out=tmp_path/'bad-replacement';assert run(p,response,evidence,out)==0
    state=get(out/'continuation-state.json')
    assert 'superseded' in state['response_holds'][key]
    assert not state['revalidated_claims']
    assert state['ordinary_workflows']==get(p/'offline-replay-state.json')['ordinary_workflows']


def test_descendant_checks_original_ancestor_and_response_evidence(tmp_path):
    p=parent(tmp_path);response,evidence=bundle(tmp_path,p);one=tmp_path/'one'
    assert run(p,response,evidence,one)==0
    key='ordinary:SYN-A:r1:supporting_evidence'
    response2,evidence2=bundle(tmp_path,one,name='two',specs=[('A',key,'merchant-payment',[])])
    two=tmp_path/'two';assert run(one,response2,evidence2,two)==0
    original=p/'sources/receipt-binder.pdf';original.write_bytes(original.read_bytes()+b'tampered')
    with pytest.raises(ValueError,match='parent evidence changed'):
        verify_continuation(two)


def test_current_corrected_revision_resumes_but_old_revision_response_is_rejected(tmp_path):
    p=parent(tmp_path,corrected=True)
    old=get(p/'offline-replay-state.json');assert old['ordinary_workflows']['SYN-A']['revision']==2
    stale,evidence=bundle(tmp_path,p)
    assert run(p,stale,evidence,tmp_path/'stale')==1
    item='correction:SYN-A:r2:expense-evidence'
    assert old['repair_queue']['items'][item]['status']=='open'
    response,evidence2=bundle(tmp_path,p,name='current',specs=[('A',item,'receipt',[]),('A',item,'merchant-payment',[])])
    out=tmp_path/'corrected';assert run(p,response,evidence2,out)==0
    state=get(out/'continuation-state.json')
    assert state['repair_queue']['items'][item]['status']=='resolved'
    assert state['ordinary_workflows']['SYN-A']['revision']==2
    assert state['ordinary_workflows']['SYN-A']['status']=='pending'
    assert state['revision_history']==old['revision_history']
    assert state['finance_reconciliations']==old['finance_reconciliations']
